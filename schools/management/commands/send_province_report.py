"""
ММО-63 I давааны өмнө аймаг/дүүргийн мэргэжилтнүүдэд (Province-ийн админууд) илгээх бүртгэлийн дэлгэрэнгүй тайлан:
холбогдсон/холбогдоогүй сургуулиуд, утсаар холбогдох эрэмбэтэй жагсаалт (Excel хавсралттай).

Холбогдсон = бүртгэгч багш эсвэл удирдлага хичээлийн жил эхэлснээс (09-01) хойш mmo.mn-д нэвтэрсэн эсвэл хандсан
(last_login эсвэл UserMeta.last_activity).
Бүртгэгч ч, удирдлага ч байхгүй сургууль холбогдоогүйд орно.

Илгээх, тест, лог нь send_province_update-тэй ижил; лог: logs/province_report_*.csv.

  python manage.py send_province_report                       # тоо, жишээ (юу ч илгээхгүй)
  python manage.py send_province_report --test-email me@example.com --province 1 28
  python manage.py send_province_report --send --add-recipient 22:name@moes.edu.mn
"""
import io
import re
from collections import Counter
from datetime import datetime

from django.core.management.base import CommandError
from django.utils import timezone

from django.db.models import Max

from emails.models import EmailEvent
from schools import enrollment
from schools.models import EnrollmentLog, School
from schools.moderator import email_problem

from .send_province_update import Command as UpdateCommand

TIERS = {
    1: ('Сурагчид хүлээж байна', 'mmo.mn-д нэвтэрч хүлээгдэж буй сурагчдын хүсэлтийг батлах'),
    2: ('Имэйл хүрэхгүй', 'Ажиллаж буй имэйл хаягийг авч registration@mmo.mn-д мэдэгдэх; нэвтэрч сурагчдаа бүртгэх'),
    3: ('Бүртгэгч багшгүй', 'Удирдлагатай ярьж бүртгэгч багш томилуулах'),
    4: ('Имэйл нээсэн, нэвтрээгүй', 'mmo.mn-д нэвтэрч сурагчдаа бүртгэхийг сануулах'),
    5: ('Имэйл хүргэгдсэн, нээгээгүй', 'Имэйлээ (Spam хавтсыг оруулаад) шалгаж, нэвтэрч сурагчдаа бүртгэхийг сануулах'),
}


def _date(dt):
    return timezone.localtime(dt).strftime('%Y-%m-%d') if dt else 'хэзээ ч'


def _dt(dt):
    return timezone.localtime(dt).strftime('%Y-%m-%d %H:%M') if dt else ''


def _activity(u):
    return getattr(getattr(u, 'data', None), 'last_activity', None) if u else None


def _last_seen(u):
    """Сүүлд нэвтэрсэн эсвэл хандсан (UserMeta.last_activity — хуучин session-оор орсныг ч тооцно) хугацааны хамгийн сүүлийнх."""
    seen = [t for t in (u.last_login, getattr(getattr(u, 'data', None), 'last_activity', None)) if t]
    return max(seen) if seen else None


def _mobile(u):
    m = getattr(getattr(u, 'data', None), 'mobile', None) if u else None
    return str(m) if m else ''


class Command(UpdateCommand):
    help = 'Аймаг/дүүргийн мэргэжилтнүүдэд бүртгэлийн дэлгэрэнгүй тайлан (утсаар холбогдох жагсаалттай) илгээнэ.'
    log_prefix = 'province_report'
    template = 'schools/emails/province/round1_report.txt'

    def handle(self, *args, **opts):
        if opts['update_post']:
            raise CommandError('--update-post нь send_province_update-д л хамаарна.')
        self.since = timezone.make_aware(datetime.combine(self._school_year_start(), datetime.min.time()))
        return super().handle(*args, **opts)

    def _school_year_start(self):
        today = timezone.localdate()
        return today.replace(month=9, day=1) if today.month >= 9 else today.replace(year=today.year - 1, month=9, day=1)

    def _email_status(self):
        """{'delivered': set, 'opened': set} — 10-01-нээс хойшх SES үйл явдал (жижиг үсгээр)."""
        if not hasattr(self, '_email_cache'):
            ev = EmailEvent.objects.filter(occurred_at__gte=self.since.replace(month=10))
            self._email_cache = {
                'delivered': {e.lower() for e in ev.filter(event_type=EmailEvent.DELIVERY).values_list('email', flat=True).distinct()},
                'opened': {e.lower() for e in ev.filter(event_type__in=[EmailEvent.OPEN, EmailEvent.CLICK])
                           .values_list('email', flat=True).distinct()},
            }
        return self._email_cache

    def _email_detail(self):
        """{email: {'delivered': (огноо, гарчиг), 'opened': огноо}} — 10-01-нээс хойшх хамгийн сүүлийн SES үйл явдал."""
        if not hasattr(self, '_email_detail_cache'):
            detail = {}
            ev = EmailEvent.objects.filter(occurred_at__gte=self.since.replace(month=10)).order_by('occurred_at')
            for email, kind, at, subject in ev.values_list('email', 'event_type', 'occurred_at', 'subject').iterator():
                d = detail.setdefault(email.lower(), {})
                if kind == EmailEvent.DELIVERY:
                    d['delivered'] = (at, subject)
                else:
                    d['opened'] = at
            self._email_detail_cache = detail
        return self._email_detail_cache

    def _last_enrollment(self):
        """{'enroll': {school_id: сүүлд сурагч бүртгэсэн (enroll/transfer)}, 'request': {school_id: сүүлийн сурагчийн хүсэлт}} — аль ч жилийнх."""
        if not hasattr(self, '_last_enrollment_cache'):
            def last(actions):
                return dict(EnrollmentLog.objects.filter(to_school__isnull=False, action__in=actions).values('to_school')
                            .annotate(last=Max('created_at')).values_list('to_school', 'last'))
            self._last_enrollment_cache = {
                'enroll': last([EnrollmentLog.Action.ENROLL, EnrollmentLog.Action.TRANSFER]),
                'request': last([EnrollmentLog.Action.REQUEST]),
            }
        return self._last_enrollment_cache

    def _province_data(self, province):
        rows = []
        this_year = self._this_year()
        mail = self._email_status()
        detail = self._email_detail()
        last_enroll = self._last_enrollment()
        for s in (School.objects.filter(province=province).exclude(name='Бусад')
                  .select_related('user__data', 'manager__data', 'group').order_by('name')):
            u = s.user if s.user and s.user.is_active else None
            m = s.manager if s.manager and s.manager.is_active else None
            staff = [x.pk for x in (s.user, s.manager) if x]
            students = s.group.user_set.exclude(pk__in=staff).count() if s.group_id else 0
            members = set(s.group.user_set.values_list('id', flat=True)) if s.group_id else set()
            new = len(this_year.get(s.id, set()) & members)
            pending = enrollment.pending_students(s).count()
            connected = any(x and _last_seen(x) and _last_seen(x) >= self.since for x in (u, m))
            emails = [x.email.lower() for x in (u, m) if x and x.email]
            if any(e in mail['opened'] for e in emails):
                mail_status = 'нээсэн'
            elif any(e in mail['delivered'] for e in emails):
                mail_status = 'хүргэгдсэн, нээгээгүй'
            elif not emails:
                mail_status = 'имэйлгүй'
            elif any(e in self.blocked for e in emails) or (u and email_problem(u, self.blocked)):
                mail_status = 'блоклогдсон / алдаатай'
            else:
                mail_status = 'хүргэлтийн бүртгэл алга'
            if connected:
                tier, problem = None, ''
            else:
                if pending:
                    tier = 1
                elif mail_status not in ('нээсэн', 'хүргэгдсэн, нээгээгүй'):
                    tier = 2
                elif not u:
                    tier = 3
                elif mail_status == 'нээсэн':
                    tier = 4
                else:
                    tier = 5
                if not u and not m:
                    problem = 'Бүртгэгч, удирдлага аль аль нь байхгүй'
                elif not u:
                    problem = 'Бүртгэгч багшгүй, удирдлага нэвтрээгүй'
                elif not u.last_login:
                    problem = 'Бүртгэгч багш системд огт нэвтрээгүй'
                else:
                    problem = 'Бүртгэгч багш 9-р сараас хойш нэвтрээгүй'
            delivered = max((detail[e]['delivered'] for e in emails if 'delivered' in detail.get(e, {})), default=None)
            opened = max((detail[e]['opened'] for e in emails if 'opened' in detail.get(e, {})), default=None)
            seen = max((_last_seen(x) for x in (u, m) if x and _last_seen(x)), default=None)
            rows.append({'last_seen': _date(seen), 'r_activity': _dt(_activity(u)), 'r_login_dt': _dt(u.last_login) if u else '',
                         'r_joined': _date(u.date_joined) if u else '', 'm_activity': _dt(_activity(m)),
                         'm_login_dt': _dt(m.last_login) if m else '', 'm_joined': _date(m.date_joined) if m else '',
                         'last_delivery': _date(delivered[0]) if delivered else '', 'last_subject': delivered[1] if delivered else '',
                         'last_open': _date(opened) if opened else '',
                         'last_enroll': _date(last_enroll['enroll'].get(s.id)),
                         'last_request': _date(last_enroll['request'].get(s.id)),
                         'name': s.name, 'id': s.id, 'connected': connected, 'tier': tier, 'problem': problem,
                         'mail': mail_status, 'students': students, 'new': new, 'pending': pending,
                         'registrar': u.get_full_name() if u else '', 'username': u.username if u else '',
                         'email': u.email if u else '', 'mobile': _mobile(u), 'last_login': _date(u.last_login) if u else '',
                         'manager': m.username if m else '', 'manager_email': m.email if m else '',
                         'manager_mobile': _mobile(m), 'manager_login': _date(m.last_login) if m else ''})
        calls = sorted((r for r in rows if not r['connected']), key=lambda r: (r['tier'], -r['pending'], -r['students'], r['name']))
        for r in calls:
            r['tier_label'] = TIERS[r['tier']][0]
            r['todo'] = TIERS[r['tier']][1]
        connected = sum(r['connected'] for r in rows)
        return {'province': province.name, 'total': len(rows), 'rows': rows, 'calls': calls,
                'connected': connected, 'unconnected': len(rows) - connected,
                'connected_pct': round(100 * connected / len(rows)) if rows else 0,
                'tiers': [(t, TIERS[t][0], n) for t, n in sorted(Counter(r['tier'] for r in calls).items())],
                'new_schools': sum(1 for r in rows if r['new']), 'new_students': sum(r['new'] for r in rows),
                'pending_students': sum(r['pending'] for r in rows), 'pending_schools': sum(1 for r in rows if r['pending']),
                'calls_pending': sum(r['pending'] for r in calls), 'no_phone': sum(1 for r in calls if not (r['mobile'] or r['manager_mobile']))}

    def _subject(self, d):
        return f'{self.olympiad["name"]} I даваа: {d["province"]} — бүртгэлийн тайлан, нэвтрээгүй {d["unconnected"]} сургууль'

    def _xlsx(self, d):
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
        bold = Font(bold=True)
        fills = {1: 'F4CCCC', 2: 'FCE5CD', 3: 'FFF2CC', 4: 'D9EAD3', 5: 'FFFFFF'}
        wb = Workbook()

        staff_cols = [
            ('Бүртгэгч багш', 'registrar', 24), ('Бүртгэгчийн утас', 'mobile', 12), ('Бүртгэгчийн имэйл', 'email', 28),
            ('Бүртгэгч сүүлд нэвтэрсэн', 'r_login_dt', 16), ('Бүртгэгч сүүлд хандсан', 'r_activity', 16),
            ('Бүртгэгч бүртгүүлсэн', 'r_joined', 12),
            ('Удирдлага', 'manager', 10), ('Удирдлагын утас', 'manager_mobile', 12), ('Удирдлагын имэйл', 'manager_email', 28),
            ('Удирдлага сүүлд нэвтэрсэн', 'm_login_dt', 16), ('Удирдлага сүүлд хандсан', 'm_activity', 16),
            ('Имэйлийн байдал', 'mail', 20), ('Сүүлд хүргэгдсэн имэйл', 'last_delivery', 12),
            ('Сүүлд хүргэгдсэн имэйлийн гарчиг', 'last_subject', 40), ('Сүүлд нээсэн', 'last_open', 12),
            ('Сүүлд сурагч бүртгэсэн (аль ч жил)', 'last_enroll', 14), ('Сүүлийн сурагчийн хүсэлт', 'last_request', 14),
        ]

        def sheet(ws, title, cols, rows, fill, freeze):
            ws.append([title])
            ws.append([])
            ws.append([c[0] for c in cols])
            for c in ws[3]:
                c.font = bold
                c.alignment = Alignment(wrap_text=True, vertical='top')
            for r in rows:
                ws.append([r.get(c[1], '') if isinstance(c[1], str) else c[1](r) for c in cols])
                color = fill(r)
                if color:
                    for c in ws[ws.max_row][:4]:
                        c.fill = PatternFill('solid', fgColor=color)
            for i, c in enumerate(cols, 1):
                ws.column_dimensions[get_column_letter(i)].width = c[2]
            ws.freeze_panes = freeze

        calls = [dict(r, no=i) for i, r in enumerate(d['calls'], 1)]
        sheet(wb.active, f'{d["province"]} — {self.today:%Y-%m-%d}-ны байдлаар. Бүртгэгч багш, удирдлага 9-р сараас хойш mmo.mn-д '
                         f'нэвтрээгүй, хандаагүй {d["unconnected"]} сургууль, чухлаас нь эрэмбэлсэн (1 = хамгийн яаралтай). '
                         f'«Сүүлд хандсан» 09-29-нөөс хойш бүртгэгддэг; хоосон бол тэр өдрөөс хойш хандаагүй.',
              [('№', 'no', 5), ('Эрэмбэ', 'tier', 7), ('Ангилал', 'tier_label', 22), ('Сургууль', 'name', 30),
               ('Асуудал', 'problem', 30), ('Хүлээгдэж буй хүсэлт', 'pending', 10), ('Бүртгэлтэй сурагч (бүх жил)', 'students', 12),
               *staff_cols, ('Юу хийлгэх', 'todo', 40), ('Залгасан огноо', '', 12), ('Үр дүн', '', 16), ('Тэмдэглэл', '', 30)],
              calls, lambda r: fills[r['tier']], 'E4')
        wb.active.title = 'Утсаар холбогдох'

        sheet(wb.create_sheet('Бүх сургууль'),
              f'{d["province"]} — сургууль бүрийн бүртгэл, {self.today:%Y-%m-%d}. «Холбогдсон» = бүртгэгч эсвэл удирдлага 9-р сараас '
              f'хойш нэвтэрсэн эсвэл хандсан. «Энэ жил бүртгэсэн» = 08-01-ээс хойш; «Бүртгэлтэй (бүх жил)» өмнөх жилүүдийнхийг оруулсан. '
              f'«Сүүлд хандсан» 09-29-нөөс хойш бүртгэгддэг.',
              [('Сургууль', 'name', 30), ('Холбогдсон', lambda r: 'тийм' if r['connected'] else 'үгүй', 11),
               ('Энэ жил бүртгэсэн', 'new', 11), ('Хүлээгдэж буй хүсэлт', 'pending', 11),
               ('Бүртгэлтэй сурагч (бүх жил)', 'students', 12), ('Нэвтрэх нэр', 'username', 14), *staff_cols],
              sorted(d['rows'], key=lambda r: (r['connected'], r['name'])),
              lambda r: None if r['connected'] else 'F4CCCC', 'B4')

        buf = io.BytesIO()
        wb.save(buf)
        name = re.sub(r'[^\w-]+', '_', d['province']).strip('_')
        return (f'{name}_burtgel_{self.today:%Y%m%d}.xlsx', buf.getvalue(),
                'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
