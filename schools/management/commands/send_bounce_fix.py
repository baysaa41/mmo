"""
Сүүлийн N хоногт имэйл нь хүргэгдээгүй (hard bounce / complaint) сурагчидтай сургуулиудын бүртгэгч
багш нарт засах заавар илгээнэ. Ихэвчлэн бүртгэгч сурагчийн имэйлийг буруу бичсэнээс үүсдэг.

Нэг сургуульд нэг л имэйл; өмнөх logs/bounce_fix_*.csv-д «sent» болсон сургуульд дахин илгээхгүй.
Бүртгэгчгүй / бүртгэгчийн имэйл хүрэхгүй сургуулиудыг алгасаж жагсаана.

  python manage.py send_bounce_fix                                  # тоо, жишээ (юу ч илгээхгүй)
  python manage.py send_bounce_fix --test-email me@example.com --school 783
  python manage.py send_bounce_fix --send
"""
import csv
import glob
import io
import os
import re
import time
from collections import defaultdict
from datetime import timedelta

from django.conf import settings
from django.contrib.auth.models import User
from django.core.mail import EmailMultiAlternatives, get_connection
from django.core.management.base import BaseCommand
from django.db.models.functions import Lower
from django.template.loader import render_to_string
from django.urls import reverse
from django.utils import timezone

from emails.html import text_to_html
from emails.models import EmailBounce
from emails.suppression import UNDELIVERABLE_BOUNCE_TYPES
from schools.management.commands.send_province_update import Command as ProvinceUpdate
from schools.models import School
from schools.moderator import NOTICE_FROM_EMAIL as FROM_EMAIL, blocked_emails, email_problem

LOG_GLOB = 'bounce_fix_*.csv'
MAX_LISTED = 40


class Command(BaseCommand):
    help = 'Имэйл нь хүргэгдээгүй сурагчидтай сургуулиудын бүртгэгч багш нарт засах заавар илгээнэ.'

    def add_arguments(self, parser):
        parser.add_argument('--days', type=int, default=30, help='Сүүлийн хэдэн хоногийн bounce')
        parser.add_argument('--school', type=int, nargs='+', help='Зөвхөн эдгээр сургууль (ID)')
        parser.add_argument('--test-email', help='Сонгосон сургуулиудын имэйлийг зөвхөн энэ хаяг руу илгээнэ')
        parser.add_argument('--send', action='store_true', help='Бодитоор илгээнэ')
        parser.add_argument('--yes', action='store_true', help='--send үед баталгаажуулалт асуухгүй')
        parser.add_argument('--rate', type=float, default=5, help='Секундэд илгээх дээд тоо')

    def handle(self, *args, **opts):
        self.today = timezone.localdate()
        self.olympiad = ProvinceUpdate._olympiad_info(self)
        self.blocked = blocked_emails()
        self.site = settings.SITE_URL.rstrip('/')
        students = self._students(opts['days'], opts['school'])
        messages, skipped = self._messages(students)
        total = sum(len(v) for v in students.values())
        self.stdout.write(self.style.MIGRATE_HEADING(
            f'Сургууль: {len(students)}, сурагч: {total} | имэйл: {len(messages)} | илгээгч: {FROM_EMAIL}'))
        for name, count, reason in skipped:
            self.stdout.write(self.style.WARNING(f'  Алгассан: {name} ({count}) — {reason}'))
        if opts['test_email']:
            for m in messages:
                self._email(m, opts['test_email'], f'[ТЕСТ → {m["email"]}] {m["subject"]}').send()
                self.stdout.write(self.style.SUCCESS(f'Тест: {m["school"]} → {opts["test_email"]}'))
        elif opts['send']:
            self._send_all(messages, opts)
        else:
            for m in messages[:1]:
                self.stdout.write('=' * 72 + f'\nTo: {m["email"]}\nSubject: {m["subject"]}\n' + '-' * 72 + f'\n{m["body"]}')
            for m in sorted(messages, key=lambda m: -m['count'])[:20]:
                self.stdout.write(f'  {m["count"]:4}  {m["school"][:40]:40} {m["email"]}')
            self.stdout.write(self.style.WARNING('Юу ч илгээгээгүй. --test-email эсвэл --send ашиглана уу.'))

    # ---------------------------------------------------------------- data

    def _students(self, days, school_ids):
        """{school: [User, ...]} — сургуулийн бүлэгт байгаа, одоогийн имэйл нь сүүлийн `days` хоногт bounce болсон сурагчид."""
        since = timezone.now() - timedelta(days=days)
        bounced = set(EmailBounce.objects.filter(bounce_type__in=UNDELIVERABLE_BOUNCE_TYPES, created_at__gte=since)
                      .annotate(e=Lower('email')).values_list('e', flat=True))
        schools = School.objects.exclude(group__isnull=True).select_related('user', 'manager')
        if school_ids:
            schools = schools.filter(id__in=school_ids)
        by_group = {s.group_id: s for s in schools}
        users = (User.objects.annotate(email_lower=Lower('email')).filter(email_lower__in=bounced, is_active=True)
                 .select_related('data__grade'))
        users = {u.pk: u for u in users}
        rows = User.groups.through.objects.filter(group_id__in=list(by_group), user_id__in=list(users))
        out = defaultdict(list)
        for gid, uid in rows.values_list('group_id', 'user_id'):
            s = by_group[gid]
            if uid not in (s.user_id, s.manager_id):
                out[s].append(users[uid])
        for s in out:
            out[s].sort(key=lambda u: (u.last_name or '', u.first_name or ''))
        return out

    # ------------------------------------------------------------- message

    def _messages(self, students):
        sent = self._already_sent()
        messages, skipped = [], []
        for s, users in students.items():
            u = s.user
            if s.id in sent:
                continue
            if u is None or not u.is_active:
                skipped.append((s.name, len(users), 'бүртгэгч багшгүй'))
                continue
            problem = email_problem(u, self.blocked)
            if problem:
                skipped.append((s.name, len(users), f'бүртгэгчийн имэйл: {problem}'))
                continue
            listed = [{'id': x.id, 'last_name': x.last_name, 'first_name': x.first_name, 'email': x.email,
                       'grade': self._grade(x)} for x in users[:MAX_LISTED]]
            ctx = {'olympiad': self.olympiad, 'from_email': FROM_EMAIL, 'school': s.name,
                   'registrar': u.get_full_name() or u.username, 'username': u.username, 'count': len(users),
                   'listed': listed, 'more': max(0, len(users) - MAX_LISTED),
                   'problems_url': self.site + reverse('school_email_problems', args=[s.id])}
            body = render_to_string('schools/emails/round1/bounce_fix.txt', ctx)
            messages.append({'school_id': s.id, 'school': s.name, 'email': u.email, 'count': len(users),
                             'subject': f'{self.olympiad["name"]} I даваа: {len(users)} сурагчийн имэйл хүргэгдэхгүй байна',
                             'body': body, 'html': text_to_html(body), 'attachment': self._xlsx(s, users)})
        return messages, skipped

    @staticmethod
    def _grade(user):
        meta = getattr(user, 'data', None)
        grade = meta.grade if meta and meta.grade_id else None
        return grade.name if grade and str(grade.name).isdigit() else ''

    def _xlsx(self, school, users):
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        wb = Workbook()
        ws = wb.active
        ws.title = 'Сурагчид'
        ws.append(['ID', 'Овог', 'Нэр', 'Анги', 'Хүргэгдээгүй имэйл', 'Зөв имэйл'])
        for c in ws[1]:
            c.font = Font(bold=True)
        ws['F1'].fill = PatternFill('solid', fgColor='FFF2CC')
        for u in users:
            ws.append([u.id, u.last_name, u.first_name, self._grade(u), u.email, ''])
        for col, w in zip('ABCDEF', (10, 20, 20, 8, 34, 34)):
            ws.column_dimensions[col].width = w
        ws.freeze_panes = 'A2'
        buf = io.BytesIO()
        wb.save(buf)
        name = re.sub(r'[^\w-]+', '_', school.name).strip('_')
        return (f'{name}_imeil_{self.today:%Y%m%d}.xlsx', buf.getvalue(),
                'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

    def _email(self, m, to, subject=None):
        msg = EmailMultiAlternatives(subject=subject or m['subject'], body=m['body'], from_email=FROM_EMAIL,
                                     to=[to], reply_to=[FROM_EMAIL])
        msg.attach_alternative(m['html'], 'text/html')
        msg.attach(*m['attachment'])
        return msg

    # ---------------------------------------------------------------- send

    def _log_dir(self):
        return os.path.join(settings.BASE_DIR, 'logs')

    def _already_sent(self):
        sent = set()
        for path in glob.glob(os.path.join(self._log_dir(), LOG_GLOB)):
            with open(path, newline='', encoding='utf-8-sig') as f:
                sent |= {int(r['school_id']) for r in csv.DictReader(f) if r['status'] == 'sent'}
        return sent

    def _send_all(self, messages, opts):
        if not messages:
            self.stdout.write('Илгээх имэйл алга.')
            return
        if not opts['yes'] and input(f'{len(messages)} имэйл илгээх үү? [yes/no]: ').strip().lower() != 'yes':
            self.stdout.write('Цуцаллаа.')
            return
        os.makedirs(self._log_dir(), exist_ok=True)
        log_path = os.path.join(self._log_dir(), f'bounce_fix_{timezone.localtime():%Y%m%d_%H%M%S}.csv')
        connection = get_connection()
        sent = failed = 0
        with open(log_path, 'w', newline='', encoding='utf-8') as f:
            log = csv.writer(f)
            log.writerow(['school_id', 'email', 'count', 'status', 'message_id', 'error', 'time'])
            for m in messages:
                msg = self._email(m, m['email'])
                msg.connection = connection
                try:
                    msg.send()
                    status = getattr(msg, 'anymail_status', None)
                    log.writerow([m['school_id'], m['email'], m['count'], 'sent',
                                  getattr(status, 'message_id', '') or '', '', timezone.now().isoformat()])
                    sent += 1
                except Exception as exc:
                    log.writerow([m['school_id'], m['email'], m['count'], 'failed', '', str(exc)[:300],
                                  timezone.now().isoformat()])
                    failed += 1
                f.flush()
                time.sleep(1 / opts['rate'] if opts['rate'] > 0 else 0)
        self.stdout.write(self.style.SUCCESS(f'Илгээсэн: {sent}, алдаа: {failed}. Лог: {log_path}'))
