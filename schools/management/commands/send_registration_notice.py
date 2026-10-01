"""
Жил бүрийн I давааны өмнө бүртгэлтэй холбоотой имэйл илгээнэ. Нэг имэйл хаяг руу нэг л имэйл
явна; хүлээн авагчийн үүргээс хамаарч хэсгүүд нэмэгдэнэ:
  - moderator: сургуулийн бүртгэгч багш (нэвтрэх, сурагч бүртгэх, хариу оруулах заавар;
               удирдлагын имэйл дутуу/ижил бол түүнийг оруулах шаардлага)
  - manager:   бүртгэгч багшаас өөр имэйлтэй сургуулийн удирдлага (багш солигдсон бол солих заавар)
  - province:  аймаг/дүүргийн админ (сургуулиудын бүртгэлд хяналт тавих, CSV хавсралттай)

Анхдагчаар юу ч илгээхгүй — хүлээн авагчдын тоо болон төрөл бүрийн жишээ имэйлийг харуулна.
  python manage.py send_registration_notice
  python manage.py send_registration_notice --test-email me@example.com
  python manage.py send_registration_notice --send
"""
import csv
import io
import os
import time
from collections import defaultdict
from datetime import datetime, time as dtime

from django.conf import settings
from django.core.mail import EmailMultiAlternatives, get_connection
from django.core.management.base import BaseCommand, CommandError
from django.template.loader import render_to_string
from django.utils import timezone

from accounts.models import Province
from olympiad.models import Olympiad, Result, SchoolYear
from schools.models import School
from schools.institutional import is_province_contact_account, is_school_manager_account
from schools.moderator import (ACCOUNT_MISSING, NOTICE_FROM_EMAIL as FROM_EMAIL, blocked_emails, email_problem,
                               manager_has_own_email, norm_email, province_admins)

ROLES = ['moderator', 'manager', 'province']
WEEKDAYS = ['Даваа', 'Мягмар', 'Лхагва', 'Пүрэв', 'Баасан', 'Бямба', 'Ням']
SUBJECTS = {  # хамгийн "өндөр" үүргээр гарчиг сонгоно
    'province': '{prefix}: сургуулиудын бүртгэлд хяналт тавих',
    'manager': '{prefix}: сургуулийн бүртгэгч багшаа шалгана уу',
    'moderator': '{prefix}: сургуулийн бүртгэл',
}


def day_genitive(day):
    """28 -> "28-ны", 21 -> "21-ний" (нэг, дөрөв, ес-өөр төгссөн өдрүүд -ний)."""
    return f'{day}-ний' if day % 10 in (1, 4, 9) else f'{day}-ны'


def place_genitive(name):
    """'Архангай аймаг' -> 'Архангай аймгийн', 'Баянзүрх дүүрэг' -> 'Баянзүрх дүүргийн'."""
    for nominative, genitive in (('аймаг', 'аймгийн'), ('дүүрэг', 'дүүргийн'), ('хот', 'хотын')):
        if name.endswith(nominative):
            return name[:-len(nominative)] + genitive
    return name


def display_name(user):
    return user.get_full_name().strip() or user.username


class Command(BaseCommand):
    help = 'I давааны өмнө бүртгэгч багш, сургуулийн удирдлага, аймгийн админуудад бүртгэлийн имэйл илгээнэ.'

    def add_arguments(self, parser):
        parser.add_argument('--school-year', type=int, help='SchoolYear ID (анхдагч: өнөөдрийг агуулсан хичээлийн жил)')
        parser.add_argument('--role', nargs='+', choices=ROLES, default=ROLES,
                            help='Зөвхөн эдгээр үүргийн хэсгийг оруулна')
        parser.add_argument('--province', type=int, nargs='+', help='Зөвхөн эдгээр аймаг/дүүрэг (ID)')
        parser.add_argument('--school', type=int, nargs='+',
                            help='Зөвхөн эдгээр сургууль (ID) болон тэдгээрийн аймаг/дүүрэг. '
                                 'Аймгийн тайлан тухайн аймгийн бүх сургуулийг хамарна.')
        parser.add_argument('--answers-deadline', help='Хариу оруулах эцсийн хугацаа, чөлөөт текст (жишээ: "11-р сарын 4")')
        parser.add_argument('--show-all', action='store_true', help='Бүх хүлээн авагчийн имэйлийг хэвлэнэ')
        parser.add_argument('--test-email', help='Хэсгүүдийн хослол бүрээс нэг жишээг энэ хаяг руу илгээнэ')
        parser.add_argument('--to', help='Бүх имэйлийг жинхэнэ хүлээн авагч руу биш энэ хаяг руу илгээнэ '
                                         '(админгүй аймгийн мэдэгдлийг ч оруулна). --send-тэй хамт.')
        parser.add_argument('--send', action='store_true', help='Бодитоор илгээнэ')
        parser.add_argument('--yes', action='store_true', help='--send үед баталгаажуулалт асуухгүй')
        parser.add_argument('--skip-log', nargs='+', default=[],
                            help='Өмнөх ажиллагааны лог(ууд) — тэнд "sent" болсон хаягууд руу дахин илгээхгүй')
        parser.add_argument('--rate', type=float, default=10, help='Секундэд илгээх дээд тоо')

    # ------------------------------------------------------------------ setup

    def handle(self, *args, **opts):
        self.school_year = self._school_year(opts['school_year'])
        self.since = timezone.make_aware(datetime.combine(self.school_year.start, dtime.min))
        self.olympiads = list(
            Olympiad.objects.filter(school_year=self.school_year, round=1)
            .select_related('level').order_by('level__name')
        )
        if not self.olympiads:
            raise CommandError(f'{self.school_year} хичээлийн жилд I давааны олимпиад алга.')
        self.base = {
            'site': settings.SITE_URL.rstrip('/'),
            'olympiad': self._olympiad_info(),
            'answers_deadline': opts['answers_deadline'],
            'guideline_url': (f"{settings.SITE_URL.rstrip('/')}/post/?id={self.school_year.guideline_post_id}"
                              if self.school_year.guideline_post_id else None),
            'from_email': FROM_EMAIL,
        }
        self.subject_prefix = self.olympiads[0].name.split(',')[0]  # "ММО-63"

        schools = School.objects.select_related(
            'province', 'user', 'user__data', 'manager').order_by('province__name', 'name')
        province_ids = opts['province']
        if opts['school']:
            province_ids = sorted(set(School.objects.filter(id__in=opts['school']).values_list('province_id', flat=True)))
        if province_ids:
            schools = schools.filter(province_id__in=province_ids)
        # Аймгийн тайланд аймгийн бүх сургууль; багш/удирдлагын хэсэгт зөвхөн сонгосон сургуулиуд
        self.province_schools = list(schools)
        self.schools = [s for s in self.province_schools if not opts['school'] or s.id in opts['school']]

        self.blocked = blocked_emails()
        self.already_sent = self._read_skip_logs(opts['skip_log'])
        self.redirect_to = opts['to']
        self.skipped = defaultdict(lambda: defaultdict(int))
        self.recipients = defaultdict(lambda: {'users': [], 'moderator': [], 'manager': [], 'province': []})

        if 'moderator' in opts['role']:
            self._collect_moderators()
        if 'manager' in opts['role']:
            self._collect_managers()
        if 'province' in opts['role']:
            self._collect_provinces(province_ids)
        messages = [self._message(email, r) for email, r in self.recipients.items()]

        self._print_summary(messages)
        self._write_report(messages)
        if opts['test_email']:
            self._send_test(messages, opts['test_email'])
        elif opts['send']:
            self._send_all(messages, opts)
        else:
            self._preview(messages, opts['show_all'])

    def _school_year(self, sy_id):
        if sy_id:
            return SchoolYear.objects.get(pk=sy_id)
        today = timezone.localdate()
        sy = SchoolYear.objects.filter(start__lte=today, end__gte=today).first()
        if not sy:
            raise CommandError('Өнөөдрийг агуулсан хичээлийн жил олдсонгүй, --school-year заана уу.')
        return sy

    def _olympiad_info(self):
        starts = {timezone.localtime(o.start_time).date() for o in self.olympiads if o.start_time}
        # Цагийг сургууль өөрөө товлох тул имэйлд зөвхөн огноо, ангиллыг харуулна
        levels = [o.level.name for o in self.olympiads]
        date_str = ', '.join(f'{d.year} оны {d.month}-р сарын {day_genitive(d.day)} өдөр ({WEEKDAYS[d.weekday()]} гараг)'
                             for d in sorted(starts))
        # "C (5-6)" -> "C (5–6)", төгсгөлд нь "анги"
        levels_display = ', '.join(name.replace('-', '–') for name in levels) + ' анги'
        return {'name': self.olympiads[0].name, 'date': date_str or 'тодорхойгүй', 'levels': levels,
                'levels_display': levels_display,
                'last_year_participants': self._last_year_participants()}

    def _last_year_participants(self):
        """Өмнөх хичээлийн жилийн I даваанд дор хаяж нэг бодлогод хариулсан сурагчдын тоо, мянгаар."""
        prev = SchoolYear.objects.filter(end__lt=self.school_year.start).order_by('-end').first()
        if not prev:
            return None
        count = Result.objects.filter(
            olympiad__school_year=prev, olympiad__round=1, answer__isnull=False,
        ).values('contestant_id').distinct().count()
        return count // 1000 if count >= 1000 else None

    def _read_skip_logs(self, paths):
        sent = set()
        for path in paths:
            with open(path, newline='', encoding='utf-8-sig') as f:
                sent |= {row['email'] for row in csv.DictReader(f) if row['status'] == 'sent'}
        return sent

    def _is_active(self, user):
        la = getattr(getattr(user, 'data', None), 'last_activity', None)
        return bool((la and la >= self.since) or (user.last_login and user.last_login >= self.since))

    def _last_seen(self, user):
        la = getattr(getattr(user, 'data', None), 'last_activity', None)
        latest = max(filter(None, [la, user.last_login]), default=None)
        return timezone.localtime(latest).strftime('%Y-%m-%d') if latest else 'хэзээ ч'

    # ------------------------------------------------------------- recipients

    def _add(self, role, user, item):
        """Хэрэглэгчийн имэйл рүү role хэсэгт item нэмнэ; илгээх боломжгүй бол алгасна."""
        email = norm_email(user.email)
        reason = email_problem(user, self.blocked)
        if not reason and email in self.already_sent:
            reason = 'өмнө нь илгээсэн'
        if reason:
            self.skipped[role][reason] += 1
            return
        entry = self.recipients[email]
        if user not in entry['users']:
            entry['users'].append(user)
        entry[role].append(item)

    def _collect_moderators(self):
        for s in self.schools:
            if not s.user:
                continue
            same_email = s.manager and not manager_has_own_email(s)
            self._add('moderator', s.user, {
                'name': s.name, 'province': s.province.name, 'username': s.user.username,
                # Удирдлага ижил имэйлтэй бол удирдлагын аккаунтын нэвтрэх нэрийг энд л мэдэгдэнэ
                'manager_username': s.manager.username if same_email and s.manager_id != s.user_id else None,
                # "Бусад" нь аймгийн админ хөтөлдөг орлуулагч сургууль — удирдлагагүй
                'needs_manager_email': s.name != 'Бусад' and not manager_has_own_email(s),
            })

    def _collect_managers(self):
        for s in self.schools:
            if s.manager and manager_has_own_email(s):
                self._add('manager', s.manager, {
                    'id': s.id, 'name': s.name, 'username': s.manager.username,
                    'moderator': f'{display_name(s.user)} ({s.user.email or "имэйлгүй"})' if s.user else None,
                    'last_seen': self._last_seen(s.user) if s.user else None,
                })

    def _email_issues(self, s):
        """Сургуулийн багш/удирдлага руу имэйл хүрэхгүй шалтгаанууд (аймгийн админд мэдэгдэнэ)."""
        if s.name == 'Бусад':  # аймгийн админ өөрөө хөтөлдөг орлуулагч сургууль
            return ''
        # "багшийн имэйл хүргэхэд алдаа гарсан", "багш, удирдлагын имэйл бүртгэгдээгүй" гэх мэт
        by_problem = defaultdict(list)
        for label, user in (('багш', s.user), ('удирдлага', s.manager)):
            problem = email_problem(user, self.blocked)
            if problem == ACCOUNT_MISSING:
                by_problem['missing'].append('бүртгэгч багшгүй' if label == 'багш' else 'удирдлагагүй')
            elif problem:
                by_problem[problem].append(label)
        issues = by_problem.pop('missing', [])
        for problem, labels in by_problem.items():
            owner = 'багш, удирдлагын' if len(labels) == 2 else {'багш': 'багшийн', 'удирдлага': 'удирдлагын'}[labels[0]]
            issues.append(f'{owner} {problem}')
        return '; '.join(issues)

    def _collect_provinces(self, province_ids):
        provinces = Province.objects.select_related('contact_person', 'registrar').order_by('name')
        if province_ids:
            provinces = provinces.filter(id__in=province_ids)
        by_province = defaultdict(list)
        for s in self.province_schools:
            by_province[s.province_id].append(s)

        self.province_summary = []
        self.school_rows = []
        for p in provinces:
            schools = by_province[p.id]
            rows = self._province_rows(p, schools)
            admins = province_admins(p, self.blocked)
            active = sum(1 for s in schools if s.user and self._is_active(s.user))
            no_mod = sum(1 for s in schools if not s.user)
            issues = [{'name': r['Сургууль'], 'issue': r['Имэйлийн асуудал']} for r in rows if r['Имэйлийн асуудал']]
            item = {'id': p.id, 'name': p.name, 'genitive': place_genitive(p.name), 'total': len(schools), 'active': active,
                    'inactive': len(schools) - active - no_mod, 'no_moderator': no_mod,
                    'issues': issues, 'rows': rows}
            for r in rows:
                r['Аймгийн админд мэдэгдсэн'] = 'тийм' if admins else 'ҮГҮЙ (админ имэйлгүй)'
            self.school_rows.extend(rows)
            self.province_summary.append({
                'Аймаг/дүүрэг': p.name, 'Сургууль': len(schools), 'Энэ жил нэвтэрсэн': active,
                'Нэвтрээгүй': item['inactive'], 'Бүртгэгч багшгүй': no_mod,
                'Имэйл хүрэхгүй сургууль': len(issues),
                'Админы имэйл': ', '.join(sorted(norm_email(u.email) for u in admins)) or '— БАЙХГҮЙ —',
            })
            if not admins:
                self.skipped['province']['имэйлтэй админгүй аймаг'] += 1
                if self.redirect_to:  # урьдчилан харахын тулд админгүй аймгийн мэдэгдлийг ч үүсгэнэ
                    entry = self.recipients[f'(админгүй) {p.name}']
                    entry['users'] += [u for u in (p.contact_person, p.registrar) if u and u not in entry['users']]
                    entry['province'].append(item)
            for u in admins:
                self._add('province', u, item)

    def _province_rows(self, province, schools):
        rows = []
        for s in schools:
            u = s.user
            rows.append({
                'Аймаг/дүүрэг': province.name,
                'Сургууль': s.name,
                'Имэйлийн асуудал': self._email_issues(s),
                'Бүртгэгч багш': display_name(u) if u else '— байхгүй —',
                'Имэйл': u.email if u else '',
                'Утас': (getattr(getattr(u, 'data', None), 'mobile', '') or '') if u else '',
                'Сүүлд нэвтэрсэн': self._last_seen(u) if u else '',
                'Энэ жил нэвтэрсэн': ('тийм' if self._is_active(u) else 'ҮГҮЙ') if u else 'ҮГҮЙ',
                'Удирдлага (менежер)': display_name(s.manager) if s.manager else '',
                'Удирдлагын имэйл': s.manager.email if s.manager else '',
                'Удирдлага өөр имэйлтэй': 'тийм' if manager_has_own_email(s) else 'ҮГҮЙ',
            })
        # Асуудалтай, дараа нь нэвтрээгүй сургуулиуд эхэндээ
        rows.sort(key=lambda r: (not r['Имэйлийн асуудал'], r['Энэ жил нэвтэрсэн'] == 'тийм', r['Сургууль']))
        return rows

    def _csv(self, rows):
        buf = io.StringIO()
        if rows:
            writer = csv.DictWriter(buf, fieldnames=list(rows[0]))
            writer.writeheader()
            writer.writerows(rows)
        return '﻿' + buf.getvalue()  # Excel-д кирилл зөв харагдахын тулд BOM

    def _message(self, email, r):
        roles = [role for role in ROLES if r[role]]
        context = {
            **self.base,
            # Удирдлагын хэсэг л байгаа бол нэрээр биш ерөнхийгөөр мэндэлнэ (автомат "Менежер ..." нэрүүд)
            'name': self._greeting_name(roles, r),
            'mod_schools': r['moderator'],
            'needs_manager_email': [s for s in r['moderator'] if s['needs_manager_email']],
            'mgr_schools': r['manager'],
            'provinces': r['province'],
            # Нэг имэйл олон аккаунтад давхцдаг тул нууц үгийг хэрэглэгчийн нэрээр сэргээхийг зөвлөнө
            'usernames': self._usernames(r),
            'roles': roles,
        }
        body = render_to_string('schools/emails/registration/notice.txt', context)
        html = render_to_string('schools/emails/registration/notice.html', context)
        top_role = next(role for role in SUBJECTS if role in roles)
        rows = [{k: v for k, v in row.items() if k != 'Аймгийн админд мэдэгдсэн'}
                for p in r['province'] for row in p['rows']]
        return {'email': email, 'users': r['users'], 'roles': roles,
                'subject': SUBJECTS[top_role].format(prefix=self.subject_prefix), 'body': body, 'html': html,
                'attachment': ('surguuliud_burtgel.csv', self._csv(rows)) if rows else None}

    def _greeting_name(self, roles, r):
        """Албан аккаунт (сургуулийн удирдлага, аймгийн удирдах ажилтан) бол нэргүйгээр "Сайн байна уу," гэж мэндэлнэ."""
        if not r['users'] or roles == ['manager']:
            return None
        user = r['users'][0]
        if is_province_contact_account(user) or is_school_manager_account(user):
            return None
        return display_name(user)

    def _usernames(self, r):
        names = [s['username'] for s in r['moderator']]
        names += [s['manager_username'] for s in r['moderator'] if s['manager_username']]
        names += [s['username'] for s in r['manager']]
        names += [u.username for u in r['users']]
        return list(dict.fromkeys(names))

    # ---------------------------------------------------------------- output

    def _write_report(self, messages):
        """Нэгдсэн тайлан: аймгууд, сургуулиуд, имэйлийн хүлээн авагчид — нэг Excel файлд."""
        from openpyxl import Workbook
        from openpyxl.utils import get_column_letter

        sheets = {
            'Аймгууд': getattr(self, 'province_summary', []),
            'Сургуулиуд': getattr(self, 'school_rows', []),
            'Имэйлүүд': [{'Имэйл': m['email'], 'Хэсгүүд': '+'.join(m['roles']),
                          'Хэрэглэгч': ', '.join(u.username for u in m['users']), 'Гарчиг': m['subject']}
                         for m in messages],
            'Алгассан': [{'Үүрэг': role, 'Шалтгаан': k, 'Тоо': v}
                         for role, d in self.skipped.items() for k, v in d.items()],
        }
        wb = Workbook()
        wb.remove(wb.active)
        for title, rows in sheets.items():
            ws = wb.create_sheet(title)
            if not rows:
                continue
            headers = list(rows[0])
            ws.append(headers)
            for row in rows:
                ws.append([row.get(h, '') for h in headers])
            ws.freeze_panes = 'A2'
            ws.auto_filter.ref = ws.dimensions
            for i, h in enumerate(headers, 1):
                width = max(len(str(h)), *(len(str(r.get(h, ''))) for r in rows[:500]))
                ws.column_dimensions[get_column_letter(i)].width = min(width + 2, 50)

        log_dir = os.path.join(settings.BASE_DIR, 'logs')
        os.makedirs(log_dir, exist_ok=True)
        path = os.path.join(log_dir, f'registration_report_{timezone.localtime():%Y%m%d_%H%M%S}.xlsx')
        wb.save(path)
        self.stdout.write(f'  Нэгдсэн тайлан: {path}')

    def _print_summary(self, messages):
        o = self.base['olympiad']
        self.stdout.write(self.style.MIGRATE_HEADING(
            f'{self.school_year} | {o["name"]} | {o["date"]} | илгээгч: {FROM_EMAIL}'))
        combos = defaultdict(int)
        for m in messages:
            combos['+'.join(m['roles'])] += 1
        self.stdout.write(f'  Нийт имэйл: {len(messages)}')
        for combo, n in sorted(combos.items(), key=lambda x: -x[1]):
            self.stdout.write(f'    {combo:28} {n}')
        for role, skipped in self.skipped.items():
            self.stdout.write(f'  Алгассан ({role}): ' + ', '.join(f'{k}: {v}' for k, v in skipped.items()))

    def _print_message(self, m):
        self.stdout.write('=' * 72)
        self.stdout.write(f'To: {m["email"]}\nFrom: {FROM_EMAIL}\nSubject: {m["subject"]}')
        if m['attachment']:
            self.stdout.write(f'Хавсралт: {m["attachment"][0]} ({m["attachment"][1].count(chr(10)) - 1} мөр)')
        self.stdout.write('-' * 72)
        self.stdout.write(m['body'])

    def _samples(self, messages):
        """Хэсгүүдийн хослол бүрээс нэг жишээ."""
        seen = {}
        for m in messages:
            seen.setdefault(tuple(m['roles']), m)
        return list(seen.values())

    def _preview(self, messages, show_all):
        for m in messages if show_all else self._samples(messages):
            self._print_message(m)
        self.stdout.write(self.style.WARNING('\nЮу ч илгээгээгүй. --test-email эсвэл --send ашиглана уу.'))

    def _email(self, m, to):
        msg = EmailMultiAlternatives(subject=m['subject'], body=m['body'], from_email=FROM_EMAIL,
                                     to=[to], reply_to=[FROM_EMAIL])
        msg.attach_alternative(m['html'], 'text/html')
        if m['attachment']:
            msg.attach(m['attachment'][0], m['attachment'][1], 'text/csv')
        return msg

    def _send_test(self, messages, to):
        for m in self._samples(messages):
            self._email({**m, 'subject': f'[ТЕСТ → {m["email"]}] {m["subject"]}'}, to).send()
            self.stdout.write(self.style.SUCCESS(f'Тест ({"+".join(m["roles"])}) илгээлээ → {to}'))

    def _send_all(self, messages, opts):
        if not messages:
            self.stdout.write('Илгээх имэйл алга.')
            return
        if not opts['yes']:
            target = f' → {self.redirect_to}' if self.redirect_to else ''
            answer = input(f'{len(messages)} имэйлийг {FROM_EMAIL} хаягаас{target} илгээх үү? [yes/no]: ')
            if answer.strip().lower() != 'yes':
                self.stdout.write('Цуцаллаа.')
                return

        log_dir = os.path.join(settings.BASE_DIR, 'logs')
        os.makedirs(log_dir, exist_ok=True)
        log_path = os.path.join(log_dir, f'registration_notice_{timezone.localtime():%Y%m%d_%H%M%S}.csv')
        interval = 1 / opts['rate'] if opts['rate'] > 0 else 0
        sent = failed = 0
        connection = get_connection()

        with open(log_path, 'w', newline='', encoding='utf-8') as f:
            log = csv.writer(f)
            log.writerow(['roles', 'email', 'user_ids', 'status', 'message_id', 'error', 'time'])
            for i, m in enumerate(messages, 1):
                if self.redirect_to:
                    msg = self._email({**m, 'subject': f'[ТЕСТ → {m["email"]}] {m["subject"]}'}, self.redirect_to)
                else:
                    msg = self._email(m, m['email'])
                msg.connection = connection
                user_ids = ' '.join(str(u.id) for u in m['users'])
                try:
                    msg.send()
                    status = getattr(msg, 'anymail_status', None)
                    # --to үед жинхэнэ хүлээн авагчид очоогүй тул --skip-log алгасахгүй байх статус
                    log.writerow(['+'.join(m['roles']), m['email'], user_ids,
                                  f'test→{self.redirect_to}' if self.redirect_to else 'sent',
                                  getattr(status, 'message_id', '') or '', '', timezone.now().isoformat()])
                    sent += 1
                except Exception as exc:
                    log.writerow(['+'.join(m['roles']), m['email'], user_ids, 'failed',
                                  '', str(exc)[:300], timezone.now().isoformat()])
                    failed += 1
                f.flush()
                if i % 50 == 0:
                    self.stdout.write(f'  {i}/{len(messages)}...')
                time.sleep(interval)

        self.stdout.write(self.style.SUCCESS(f'Илгээсэн: {sent}, алдаа: {failed}. Лог: {log_path}'))
        if failed:
            self.stdout.write(f'Алдаатайг дахин илгээх: --send --skip-log {log_path}')
